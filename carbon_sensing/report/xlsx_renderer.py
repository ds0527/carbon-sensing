"""articles.xlsx. 전체 문서를 필터·정렬 가능한 표로 떨어뜨린다."""

from __future__ import annotations

import logging
from pathlib import Path
from typing import Any

from ..daterange import KST
from ..models import Document

log = logging.getLogger(__name__)

HEADERS = [
    "순위",
    "점수",
    "카테고리",
    "대표",
    "제목",
    "출처",
    "유형",
    "지역",
    "게재일",
    "중복보도",
    "본문확보",
    "수집기",
    "제외사유",
    "URL",
]


def _row(doc: Document, rank: int, is_representative: bool) -> list[Any]:
    published = (
        doc.published_at.astimezone(KST).strftime("%Y-%m-%d")
        if doc.published_at
        else "일자미확인"
    )
    return [
        rank,
        doc.score,
        doc.category,
        "O" if is_representative else "",
        doc.title,
        doc.source_name,
        doc.doc_type,
        doc.region,
        published,
        doc.duplicate_count,
        "O" if doc.has_body else "",
        doc.collector,
        doc.score_detail.get("off_topic") or "",
        doc.url,
    ]


def write_xlsx(result: Any, path: Path) -> Path | None:
    """openpyxl이 없으면 건너뛰고 None을 반환한다."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
        from openpyxl.utils import get_column_letter
    except ImportError:
        log.warning("openpyxl이 없어 articles.xlsx를 건너뜁니다.")
        return None

    representative_ids = {d.id for d in result.representative}

    wb = Workbook()
    ws = wb.active
    ws.title = "문서목록"
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    rank = 0
    for doc in result.documents:
        rank += 1
        ws.append(_row(doc, rank, doc.id in representative_ids))
    for doc in result.date_unknown:
        rank += 1
        ws.append(_row(doc, rank, doc.id in representative_ids))

    # 제목·URL은 넓게, 나머지는 내용에 맞춰 폭을 준다.
    widths = [6, 8, 10, 6, 60, 18, 10, 8, 12, 10, 10, 12, 16, 60]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{ws.max_row}"

    # 두 번째 시트: 실행 조건과 수집 현황
    meta = wb.create_sheet("실행조건")
    meta.append(["항목", "값"])
    for cell in meta[1]:
        cell.font = Font(bold=True)
    ctx = result.context
    rows = [
        ("기간", result.period.label()),
        ("기간 일수", result.period.days),
        ("기본 주제", ctx.base_topic),
        ("추가 주제", ", ".join(ctx.extra_topics) or "없음"),
        ("결합 모드", ctx.mode.upper()),
        ("제외어", ", ".join(ctx.excludes) or "없음"),
        ("질의어", " | ".join(ctx.queries)),
        ("원시 수집", result.stats.get("raw", 0)),
        ("중복제거 후", result.stats.get("deduped", 0)),
        ("기간 내", result.stats.get("in_period", 0)),
        ("관련 문서(임계값 이상)", len(result.visible)),
        ("대표 기사", len(result.representative)),
        ("LLM 토큰", result.usage.get("total_tokens", 0) if result.usage else 0),
        (
            "LLM 추정비용(USD)",
            result.usage.get("estimated_cost_usd", 0) if result.usage else 0,
        ),
    ]
    for key, value in rows:
        meta.append([key, value])
    meta.column_dimensions["A"].width = 24
    meta.column_dimensions["B"].width = 90

    collectors = wb.create_sheet("수집기")
    collectors.append(["수집기", "상태", "원시건수", "비고"])
    for cell in collectors[1]:
        cell.font = Font(bold=True)
    for cr in result.collector_results:
        note = cr.skipped_reason or ("; ".join(cr.errors) if cr.errors else "")
        collectors.append([cr.collector, cr.status, len(cr.documents), note])
    collectors.column_dimensions["A"].width = 14
    collectors.column_dimensions["B"].width = 10
    collectors.column_dimensions["C"].width = 10
    collectors.column_dimensions["D"].width = 90

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
